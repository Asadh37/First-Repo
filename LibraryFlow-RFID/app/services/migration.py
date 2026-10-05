from __future__ import annotations
import json
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from ..models import AuditEvent, Book, Loan, MigrationBatch, MigrationStaging
from .audit import record

REQUIRED_COLUMNS = {"accession_no", "title"}
ALIASES = {"accession": "accession_no", "accession number": "accession_no", "accession_no": "accession_no", "title": "title", "book title": "title", "author": "author", "isbn": "isbn", "shelf": "shelf", "location": "shelf"}

def _normalise_header(value):
    key = str(value or "").strip().lower().replace("-", "_")
    return ALIASES.get(key, key.replace(" ", "_"))

def stage_xlsx(db: Session, file_path: str, source_name: str):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try: header = [_normalise_header(x) for x in next(rows)]
    except StopIteration: raise ValueError("Spreadsheet is empty")
    if not REQUIRED_COLUMNS.issubset(set(header)):
        raise ValueError("Spreadsheet must include accession_no and title columns")
    batch = MigrationBatch(source_name=source_name, status="staging")
    db.add(batch); db.flush()
    seen = set(); existing = set(db.scalars(select(Book.accession_no)).all())
    valid = rejected = duplicates = source_count = 0
    errors = []
    for row_num, values in enumerate(rows, start=2):
        if not any(v is not None and str(v).strip() for v in values): continue
        source_count += 1
        data = {h: (str(values[i]).strip() if i < len(values) and values[i] is not None else "") for i,h in enumerate(header)}
        accession = data.get("accession_no", "").strip(); title = data.get("title", "").strip()
        errs = []
        if not accession: errs.append("missing accession_no")
        if not title: errs.append("missing title")
        status = "valid"
        if errs: status = "rejected"; rejected += 1
        elif accession in seen or accession in existing: status = "duplicate"; duplicates += 1; errs.append("accession_no already exists in batch or target")
        else: seen.add(accession); valid += 1
        db.add(MigrationStaging(batch_id=batch.id, row_number=row_num, accession_no=accession, title=title, author=data.get("author", ""), isbn=data.get("isbn", ""), shelf=data.get("shelf", "A-01") or "A-01", status=status, errors="; ".join(errs)))
        if errs: errors.append({"row": row_num, "accession_no": accession, "errors": errs})
    batch.source_rows=source_count; batch.valid_rows=valid; batch.rejected_rows=rejected; batch.duplicate_rows=duplicates; batch.status="staged"
    batch.report_json=json.dumps({"errors":errors},ensure_ascii=False)
    record(db,"migration.staged","migration_batch",batch.id,{"source_rows":source_count,"valid_rows":valid,"rejected_rows":rejected,"duplicates":duplicates})
    db.commit(); wb.close()
    return batch

def commit_batch(db: Session, batch_id: int):
    batch = db.get(MigrationBatch, batch_id)
    if not batch: raise ValueError("Migration batch not found")
    if batch.status != "staged": raise ValueError(f"Batch must be staged; current status: {batch.status}")
    rows = db.scalars(select(MigrationStaging).where(MigrationStaging.batch_id==batch_id, MigrationStaging.status=="valid")).all()
    created=[]
    try:
        existing=set(db.scalars(select(Book.accession_no)).all())
        for r in rows:
            if r.accession_no in existing:
                r.status="duplicate"; r.errors="accession_no became duplicate before commit"; batch.duplicate_rows += 1; batch.valid_rows -= 1
                continue
            db.add(Book(accession_no=r.accession_no,title=r.title,author=r.author,isbn=r.isbn,shelf=r.shelf))
            created.append(r.accession_no); existing.add(r.accession_no)
        batch.status="committed"
        record(db,"migration.committed","migration_batch",batch.id,{"created_count":len(created),"accession_numbers":created})
        db.commit()
        return {"batch_id":batch.id,"status":batch.status,"created_count":len(created),"created_accessions":created,"source_rows":batch.source_rows,"rejected_rows":batch.rejected_rows,"duplicate_rows":batch.duplicate_rows,"reconciled":len(created)+batch.rejected_rows+batch.duplicate_rows==batch.source_rows}
    except Exception:
        db.rollback(); raise

def rollback_batch(db: Session, batch_id: int):
    batch=db.get(MigrationBatch,batch_id)
    if not batch: raise ValueError("Migration batch not found")
    if batch.status != "committed": raise ValueError("Only a committed batch can be rolled back")
    audit = db.execute(select(AuditEvent).where(AuditEvent.action=="migration.committed", AuditEvent.entity_id==str(batch_id)).order_by(AuditEvent.id.desc())).scalars().first()
    if not audit: raise ValueError("Missing migration commit audit evidence; safe rollback refused")
    payload=json.loads(audit.details)
    accessions=payload.get("accession_numbers",[])
    for accession in accessions:
        book=db.scalar(select(Book).where(Book.accession_no==accession))
        if book:
            if book.rfid_tag: raise ValueError(f"Safe rollback refused: {accession} has an RFID tag")
            if db.scalar(select(Loan.id).where(Loan.book_id==book.id).limit(1)):
                raise ValueError(f"Safe rollback refused: {accession} has circulation history")
            db.delete(book)
    batch.status="rolled_back"
    record(db,"migration.rolled_back","migration_batch",batch.id,{"removed_count":len(accessions),"accession_numbers":accessions})
    db.commit()
    return {"batch_id":batch.id,"status":batch.status,"removed_count":len(accessions),"accession_numbers":accessions}
